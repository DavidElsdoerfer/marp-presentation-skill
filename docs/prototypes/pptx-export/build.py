#!/usr/bin/env python3
"""Prototyp: JSON-Modell (model.cjs) -> PPTX auf den Layouts einer Vorlage (python-pptx).

  python build.py model.json brand.json vorlage.pptx ausgabe.pptx

Benötigt python-pptx (nicht Teil des Skills; optional, in eigener Umgebung installieren). Die Vorlage darf Folien enthalten, sie werden nicht
entfernt (Prototyp): am besten die leere Vorlage ohne Folien verwenden (siehe README: pptx_showcase-Variante ohne Folien).
brand.json: "layouts" (Klasse -> Layoutname) und "custom_layouts", wie sie `marp-deck brand import` erzeugt.
"""
import json, sys
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.util import Emu

model, brand = json.load(open(sys.argv[1])), json.load(open(sys.argv[2]))
prs = Presentation(sys.argv[3])

# Achtung: prs.slide_layouts enthält nur die Layouts des ERSTEN Masters. Alle Master durchlaufen.
layouts = {l.name: l for m in prs.slide_masters for l in m.slide_layouts}
cls_to_layout = {**brand["layouts"], **brand.get("custom_layouts", {})}
cls_to_layout = {k: v.replace(" (wie Titelfolie)", "") for k, v in cls_to_layout.items()}   # Alias-Hinweis des Imports entfernen


def text_of(runs):
    return "".join(r.get("t", "") for r in runs if "t" in r)


def fill_runs(par, runs):
    for r in runs:
        if "t" not in r:
            continue
        run = par.add_run(); run.text = r["t"]
        if r.get("b"): run.font.bold = True
        if r.get("i"): run.font.italic = True
        if r.get("code"): run.font.name = "Courier New"
        if r.get("link"): run.hyperlink.address = r["link"]


def fill_body(tf, blocks):
    first = True
    for b in blocks:
        if b["type"] in ("paragraph", "quote"):
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            fill_runs(p, b["runs"])
        elif b["type"] == "list":
            for it in b["items"]:
                p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
                p.level = min(it["level"], 8)           # Aufzählungszeichen und Einzug kommen aus dem Master (bodyStyle)
                fill_runs(p, it["runs"])


def placeholders(slide, *types):
    return [s for s in slide.placeholders if s.placeholder_format.type in types]


for s in model["slides"]:
    cls = s["directives"].get("class") or "content"
    layout = layouts[cls_to_layout.get(cls) or cls_to_layout["content"]]      # Prototyp: Rückfall auf "content"; im Produkt: harter Abbruch
    slide = prs.slides.add_slide(layout)                                       # kopiert Titel-/Inhaltsplatzhalter, NICHT Datum/Fuß/Nummer
    heads = [b for b in s["blocks"] if b["type"] == "heading"]
    h1 = next((h for h in heads if h["level"] == 1), None)
    h2 = next((h for h in heads if h["level"] == 2), None)
    title = placeholders(slide, PH.TITLE, PH.CENTER_TITLE)
    if title and h1:
        title[0].text_frame.text = text_of(h1["runs"])
    others = [b for b in s["blocks"] if b["type"] != "heading"]
    # Inhaltsplatzhalter in Leserichtung (erst oben nach unten, dann links nach rechts) = Reihenfolge der Slots
    bodies = sorted(placeholders(slide, PH.BODY, PH.OBJECT, PH.SUBTITLE), key=lambda p: (round(p.top / 300000), p.left))
    sub = [p for p in bodies if p.placeholder_format.type == PH.SUBTITLE]
    if sub and h2:
        sub[0].text_frame.text = text_of(h2["runs"]); bodies = [p for p in bodies if p not in sub]
    elif cls in ("section", "closing") and h2 and bodies:                      # Abschnittsnummer/Kontakt steht oft im Textfeld
        bodies[0].text_frame.text = text_of(h2["runs"]); bodies = bodies[1:]
    slots = {}
    for b in others:
        slots.setdefault(b["slot"], []).append(b)
    if None in slots and len(slots) == 1:                                      # Fließtext ohne Slots
        if bodies:
            body = bodies[0]
            fill_body(body.text_frame, [b for b in slots[None] if b["type"] != "table"])
            for t in (b for b in slots[None] if b["type"] == "table"):
                rows, cols = len(t["rows"]), len(t["rows"][0])
                gf = slide.shapes.add_table(rows, cols, body.left, body.top + Emu(900000 if len(slots[None]) > 1 else 0), body.width, Emu(400000 * rows))
                for r, row in enumerate(t["rows"]):
                    for c, cell in enumerate(row):
                        gf.table.cell(r, c).text = text_of(cell)               # Tabellenstil (Standard: Medium Style 2 – Accent 1) kommt aus dem Theme
                if all(b["type"] == "table" for b in slots[None]):
                    body._element.getparent().remove(body._element)
    else:
        for n, key in enumerate(sorted(k for k in slots if k is not None)):     # n-ter Slot -> n-ter Inhaltsplatzhalter
            if n < len(bodies):
                fill_body(bodies[n].text_frame, slots[key])
    for bl in others:
        if bl["type"] == "html":
            print(f"  Folie {s['index'] + 1}: HTML-Komponente nicht abbildbar: {bl['classes']}")   # Produkt: Bild der Folie als Rückfall
    for p in list(slide.placeholders):                                         # leere Platzhalter entfernen (sonst "Text hier eingeben")
        if p.has_text_frame and not p.text_frame.text.strip():
            p._element.getparent().remove(p._element)
    if s["notes"]:
        slide.notes_slide.notes_text_frame.text = "\n".join(s["notes"])         # legt bei Bedarf einen Notizmaster an
prs.save(sys.argv[4])
print(f"gespeichert: {len(prs.slides)} Folien -> {sys.argv[4]}")
